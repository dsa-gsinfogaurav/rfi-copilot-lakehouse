"""Build an editable PowerPoint from deck/RFI_Copilot_Business_Case.md (slides separated by ---).
Native text boxes and tables (no images), 16:9. Requires python-pptx."""
import math, os, re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "deck")
SRC = os.path.join(ROOT, "RFI_Copilot_Business_Case.md")
OUT = os.path.join(ROOT, "RFI_Copilot_Business_Case.pptx")
INK, RED, STEEL, MUT, LINE, ZEBRA, PALE = (RGBColor(0x14, 0x20, 0x2B), RGBColor(0xC2, 0x1E, 0x2E), RGBColor(0x3E, 0x5C, 0x76),
                                          RGBColor(0x8A, 0x97, 0xA3), RGBColor(0xD9, 0xDE, 0xE3), RGBColor(0xF5, 0xF7, 0xF9), RGBColor(0xC6, 0xD2, 0xDA))
FONT = "Helvetica Neue"
W, H, LM, CW = 13.333, 7.5, 0.75, 13.333 - 1.5

def runs(par, text, size, color=INK, bold=False):
    """Write markdown inline (**bold**, *italic*, `code`) into a paragraph as runs."""
    for tok in re.split(r"(\*\*.+?\*\*|`.+?`|(?<![\w*])\*(?!\s).+?(?<!\s)\*(?![\w*]))", text):
        if not tok: continue
        r = par.add_run(); f = r.font; f.size = Pt(size); f.color.rgb = color; f.name = FONT; f.bold = bold
        if tok.startswith("**"): r.text = tok[2:-2]; f.bold = True
        elif tok.startswith("`"): r.text = tok[1:-1]; f.name = "Menlo"; f.size = Pt(size - 1)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2: r.text = tok[1:-1]; f.italic = True
        else: r.text = tok

def plain(t): return re.sub(r"\*\*|`|(?<![\w*])\*(?!\s)|(?<!\s)\*(?![\w*])", "", t)
def est_lines(text, width_in, size):  # rough wrap estimate for layout
    cpl = max(10, int(width_in * 72 / (0.5 * size)))
    return max(1, math.ceil(len(plain(text)) / cpl))

def parse(lines):
    blocks, i = [], 0
    while i < len(lines):
        l = lines[i]
        if not l.strip(): i += 1; continue
        if l.startswith("# "): blocks.append(("h1", l[2:])); i += 1; continue
        if l.startswith("## "): blocks.append(("h2", l[3:])); i += 1; continue
        if l.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells): rows.append(cells)
                i += 1
            blocks.append(("table", rows)); continue
        if re.match(r"^(- |\d+\. )", l):
            num = bool(re.match(r"^\d+\. ", l)); items = []
            while i < len(lines) and re.match(r"^(- |\d+\. )", lines[i]):
                items.append(re.sub(r"^(- |\d+\. )", "", lines[i])); i += 1
            blocks.append(("ol" if num else "ul", items)); continue
        if l.startswith("> "): blocks.append(("quote", l[2:])); i += 1; continue
        blocks.append(("p", l)); i += 1
    return blocks

def textbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame
    tf.word_wrap = True; tf.margin_left = tf.margin_right = Inches(0); tf.margin_top = tf.margin_bottom = Inches(0.02)
    return tb, tf

def add_table(slide, rows, y, size):
    ncol = len(rows[0])
    widths_raw = [max(len(plain(r[c])) if c < len(r) else 0 for r in rows) for c in range(ncol)]
    widths_raw = [min(max(w, 8), 70) for w in widths_raw]
    tot = sum(widths_raw); widths = [CW * w / tot for w in widths_raw]
    heights = [max(est_lines(r[c] if c < len(r) else "", widths[c] - 0.18, size) for c in range(ncol)) * size * 1.25 / 72 + 0.12 for r in rows]
    shp = slide.shapes.add_table(len(rows), ncol, Inches(LM), Inches(y), Inches(CW), Inches(sum(heights)))
    tbl = shp.table
    tbl.first_row = True
    for c, wv in enumerate(widths): tbl.columns[c].width = Inches(wv)
    for ri, r in enumerate(rows):
        tbl.rows[ri].height = Inches(heights[ri])
        for c in range(ncol):
            cell = tbl.cell(ri, c); cell.margin_left = cell.margin_right = Inches(0.08); cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.TOP
            cell.fill.solid(); cell.fill.fore_color.rgb = INK if ri == 0 else (ZEBRA if ri % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
            tf = cell.text_frame; tf.word_wrap = True; p = tf.paragraphs[0]
            runs(p, r[c] if c < len(r) else "", size, RGBColor(0xFF, 0xFF, 0xFF) if ri == 0 else INK, bold=(ri == 0))
    return sum(heights)

def build():
    md = open(SRC).read()
    slides = [s.strip().splitlines() for s in re.split(r"\n---\n", md)]
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    for n, lines in enumerate(slides, 1):
        s = prs.slides.add_slide(blank); blocks = parse(lines)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12)); bar.fill.solid(); bar.fill.fore_color.rgb = RED; bar.line.fill.background()
        if n == 1:  # dark title slide
            bg = s.background.fill; bg.solid(); bg.fore_color.rgb = INK
            y = 1.7
            for kind, val in blocks:
                if kind == "h1":
                    _, tf = textbox(s, LM, y, CW, 1.1); runs(tf.paragraphs[0], val, 44, RGBColor(0xFF, 0xFF, 0xFF), bold=True); y += 1.25
                else:
                    h = est_lines(val, 10.5, 17) * 17 * 1.35 / 72 + 0.1
                    _, tf = textbox(s, LM, y, 10.5, h); runs(tf.paragraphs[0], val, 17, PALE); y += h + 0.28
            continue
        y = 0.55
        for kind, val in blocks:
            if kind in ("h1", "h2"):
                _, tf = textbox(s, LM, y, CW, 0.7); runs(tf.paragraphs[0], val, 26, INK, bold=True); y += 0.8
            elif kind == "table":
                dense = sum(len(plain(c)) for r in val for c in r) > 900
                y += add_table(s, val, y, 10.5 if dense else 11.5) + 0.2
            elif kind in ("ul", "ol"):
                size = 13 if sum(len(x) for x in val) > 900 else 14
                h = sum(est_lines(x, CW - 0.4, size) for x in val) * size * 1.32 / 72 + len(val) * 0.06
                _, tf = textbox(s, LM, y, CW, h)
                for j, item in enumerate(val):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph(); p.space_after = Pt(4)
                    runs(p, ("•  " if kind == "ul" else f"{j + 1}.  ") + item, size)
                y += h + 0.15
            elif kind == "quote":
                h = est_lines(val, CW - 0.4, 12) * 12 * 1.35 / 72 + 0.2
                box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(LM), Inches(y), Inches(CW), Inches(h)); box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xEE, 0xF2, 0xF5); box.line.fill.background()
                tf = box.text_frame; tf.margin_left = Inches(0.18); tf.word_wrap = True; runs(tf.paragraphs[0], val, 12, STEEL); y += h + 0.15
            else:
                h = est_lines(val, CW, 14) * 14 * 1.35 / 72 + 0.05
                _, tf = textbox(s, LM, y, CW, h); runs(tf.paragraphs[0], val, 14); y += h + 0.14
        _, tf = textbox(s, LM, H - 0.42, CW - 0.6, 0.25); runs(tf.paragraphs[0], "RFI Copilot · Summit Ridge Builders (synthetic data) · Databricks", 9, MUT)
        _, tf = textbox(s, W - LM - 0.5, H - 0.42, 0.5, 0.25); runs(tf.paragraphs[0], str(n), 9, MUT)
        if y > H - 0.5: print(f"WARNING slide {n} may overflow (content ends at {y:.2f} in)")
    prs.save(OUT); print("wrote", OUT, len(slides), "slides")

if __name__ == "__main__":
    build()
